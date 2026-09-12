"""Generate Persian Word coding tutorial."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

OUTPUT = "/home/z/my-project/download/CODING_TUTORIAL_FA.docx"
doc = Document()

for section in doc.sections:
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    sectPr = section._sectPr
    sectPr.append(OxmlElement("w:bidi"))

sn = doc.styles["Normal"]
sn.font.name = "Calibri"
sn.font.size = Pt(11)
sn.paragraph_format.line_spacing = 1.5
sn.paragraph_format.space_after = Pt(6)

for i, sz in [(1,18), (2,15), (3,13)]:
    st = doc.styles["Heading %d" % i]
    st.font.name = "Calibri"
    st.font.size = Pt(sz)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0x00, 0x66, 0xCC)

def rtl(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(11)
    pPr = p._p.get_or_add_pPr()
    pPr.append(OxmlElement("w:bidi"))
    return p

def code(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    r = p.add_run(text)
    r.font.name = "Consolas"
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x00, 0x80, 0x00)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F0F0F0")
    shd.set(qn("w:val"), "clear")
    p._p.get_or_add_pPr().append(shd)
    return p

def bullet(text):
    p = doc.add_paragraph(text, style="List Bullet")
    pPr = p._p.get_or_add_pPr()
    pPr.append(OxmlElement("w:bidi"))
    return p

# COVER
for _ in range(6): doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("آموزش کدنویسی سایت شخصی")
r.font.size = Pt(28)
r.font.bold = True
r.font.color.rgb = RGBColor(0x00, 0x66, 0xCC)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("از صفر تا صد — به زبان فارسی با مثال‌های فراوان")
r.font.size = Pt(14)
r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
doc.add_page_break()

# TOC
doc.add_heading("فهرست مطالب", level=1)
chapters = [
    "فصل ۱: نصب با Docker",
    "فصل ۲: ساختار فایل‌ها",
    "فصل ۳: مدیریت محتوا از پنل",
    "فصل ۴: TypeScript با مثال",
    "فصل ۵: React — State و Effect",
    "فصل ۶: Canvas — رسم موج",
    "فصل ۷: API — بک‌اند",
    "فصل ۸: دیتابیس Prisma",
    "فصل ۹: ساخت بخش جدید",
    "فصل ۱۰: تم‌ساز رنگی",
    "فصل ۱۱: ایمیل و Bale",
    "فصل ۱۲: ضد کپی",
    "فصل ۱۳: دستورات Docker",
    "فصل ۱۴: رفع اشکال",
]
for ch in chapters:
    rtl(ch)
doc.add_page_break()

# CH1
doc.add_heading("فصل ۱: نصب با Docker", level=1)
rtl("راحت‌ترین راه: استفاده از Docker. فقط Docker روی سیستم نصب باشه.")
code("curl -fsSL https://get.docker.com | sh\nsudo systemctl enable docker")
rtl("اجرای سایت:")
code("unzip personal-site-v7.zip\ncd personal-site-v7\ndocker-compose up -d\n# سایت روی پورت 3000")
rtl("Docker خودش همه‌چیز رو نصب می‌کنه: Node.js، Bun، پکیج‌ها، دیتابیس، محتوای اولیه.")

# CH2
doc.add_heading("فصل ۲: ساختار فایل‌ها", level=1)
code("personal-site-v7/\n  Dockerfile           تنظیمات Docker\n  docker-compose.yml   اجرای خودکار\n  prisma/schema.prisma  ساختار دیتابیس\n  src/app/page.tsx      صفحه اصلی\n  src/app/personal.css  رنگ و فونت\n  src/app/api/          بک‌اند\n  src/components/       کامپوننت‌ها\n  src/lib/content.ts    محتوای پایه\n  scripts/              اسکریپت‌های seed")
rtl("جدول: کجا چی رو تغییر بدم؟")
changes = [
    ("نام من", "src/lib/content.ts"),
    ("رنگ‌ها", "src/app/personal.css یا پنل themes"),
    ("متن بخش‌ها", "پنل ادمین (بدون کد)"),
    ("شخصیت بات", "پنل ادمین (بدون کد)"),
    ("تم جدید", "پنل ادمین themes (بدون کد)"),
    ("بخش جدید", "src/app/page.tsx"),
    ("مدل دیتابیس", "prisma/schema.prisma"),
]
tbl = doc.add_table(rows=len(changes)+1, cols=2)
tbl.style = "Table Grid"
tbl.rows[0].cells[0].text = "می‌خوام تغییر بدم"
tbl.rows[0].cells[1].text = "کجا"
for i, (a, b) in enumerate(changes):
    tbl.rows[i+1].cells[0].text = a
    tbl.rows[i+1].cells[1].text = b

# CH3
doc.add_heading("فصل ۳: مدیریت محتوا از پنل", level=1)
rtl("همه‌ی محتوا از پنل ادمین قابل مدیریت است.")
bullet("آدرس: http://localhost:3000/#admin")
bullet("یا: Ctrl+Shift+A")
bullet("رمز: admin123")
rtl("تب‌های پنل: contact messages / AI chat logs / content / themes / settings")
rtl("مثال: اضافه‌کردن آموزش:")
code('پنل ad content Tutorials + add\nTitle: Python Tutorial 5\nEmbed URL: https://www.aparat.com/video/video/embed/videohash/XXXX/vframe\nDuration: 15:00\nLevel: Beginner\nSave')
rtl("مثال: وارد کردن ۱۰۰ آموزش آپارات:")
code('[\n  {\n    "titleEn": "آموزش پایتون ۱",\n    "embedUrl": "https://www.aparat.com/video/video/embed/videohash/ABC123/vframe",\n    "duration": "10:00",\n    "levelEn": "Beginner",\n    "descEn": "مبانی پایتون"\n  }\n]\n\nدکمه import بزن — همه یکجا اضافه می‌شن')

# CH4
doc.add_heading("فصل ۴: TypeScript با مثال", level=1)
rtl("TypeScript = جاوااسکریپت + نوع‌دهی.")
code("let name: string = " + chr(34) + "YourName" + chr(34) + ";       // رشته\nlet age: number = 30;              // عدد\nlet active: boolean = true;     // درست/غلط\nlet items: string[] = [" + chr(34) + "a" + chr(34) + ", " + chr(34) + "b" + chr(34) + "]; // آرایه")
rtl("تابع با نوع:")
code("function greet(name: string): string {\n  return " + chr(34) + "سلام " + chr(34) + " + name;\n}\nlet msg = greet(" + chr(34) + "سعید" + chr(34) + ");  // سلام سعید")
rtl("interface:")
code("interface Book {\n  id: string;\n  title: string;\n  visible: boolean;\n}\nconst b: Book = { id: " + chr(34) + "b1" + chr(34) + ", title: " + chr(34) + "کتاب" + chr(34) + ", visible: true };")

# CH5
doc.add_heading("فصل ۵: React — State و Effect", level=1)
rtl("کامپوننت یه تابع است که HTML برمی‌گردونه.")
code("function Greeting({ name }) {\n  return <h1>سلام {name}</h1>;\n}\n<Greeting name=" + chr(34) + "سعید" + chr(34) + " />")
rtl("useState — متغیر حالت:")
code("function Counter() {\n  const [count, setCount] = useState(0);\n  return <button onClick={() => setCount(count + 1)}>کلیک: {count}</button>;\n}")
rtl("useEffect — اجرا هنگام لود:")
code("function BookList() {\n  const [books, setBooks] = useState([]);\n  useEffect(() => {\n    fetch('/api/content').then(r => r.json()).then(d => setBooks(d.books));\n  }, []);\n  return books.map(b => <li key={b.id}>{b.titleEn}</li>);\n}")

# CH6
doc.add_heading("فصل ۶: Canvas — رسم موج", level=1)
rtl("اسیلوسکوپ با HTML5 Canvas ساخته شده.")
code("const ctx = canvas.getContext('2d');\nctx.fillStyle = '#000';\nctx.fillRect(0, 0, w, h);\nctx.strokeStyle = '#00ff41';\nctx.beginPath();\nfor (let x = 0; x < w; x++) {\n  const y = h/2 + Math.sin(x * 0.05) * 50;\n  if (x === 0) ctx.moveTo(x, y); else ctx.lineTo(x, y);\n}\nctx.stroke();")
rtl("انیمیشن:")
code("function animate() {\n  requestAnimationFrame(animate);\n  time += 0.05;\n  ctx.fillRect(0,0,w,h);  // پاک‌کردن\n  ctx.beginPath();\n  for (let x = 0; x < w; x++) {\n    const y = h/2 + Math.sin(x*0.05 + time) * 50;\n    ctx.lineTo(x, y);\n  }\n  ctx.stroke();\n}\nanimate();")

# CH7
doc.add_heading("فصل ۷: API — بک‌اند", level=1)
rtl("هر فایل route.ts توی src/app/api/ یه API است.")
code("import { NextResponse } from 'next/server';\n\nexport async function GET() {\n  return NextResponse.json({ msg: 'سلام' });\n}\n// fetch('/api/hello') => { msg: 'سلام' }")
rtl("POST با پارامتر:")
code("export async function POST(req) {\n  const body = await req.json();\n  return NextResponse.json({ greeting: 'سلام ' + body.name });\n}")
rtl("API با دیتابیس:")
code("import { db } from '@/lib/db';\n\nexport async function GET() {\n  const books = await db.book.findMany({ where: { visible: true } });\n  return NextResponse.json({ books });\n}")

# CH8
doc.add_heading("فصل ۸: دیتابیس Prisma", level=1)
rtl("فایل prisma/schema.prisma ساختار جداول رو تعریف می‌کنه:")
code("model Book {\n  id      String  @id @default(cuid())\n  titleEn String\n  year    String\n  visible Boolean @default(true)\n  order   Int    @default(0)\n  createdAt DateTime @default(now())\n  @@index([visible, order])\n}")
rtl("استفاده تو کد:")
code("// خواندن\nconst books = await db.book.findMany({ where: { visible: true } });\n// ایجاد\nconst b = await db.book.create({ data: { titleEn: 'کتاب', year: '2025' } });\n// حذف\nawait db.book.delete({ where: { id: 'xxx' } });")
rtl("تغییر مدل:")
code("bunx prisma db push --accept-data-loss")

# CH9
doc.add_heading("فصل ۹: ساخت بخش جدید", level=1)
rtl("مثال: اضافه‌کردن بخش پروژه‌ها:")
code("۱. مدل دیتابیس اضافه کن (schema.prisma):\n   model Project { id String @id, titleEn String, visible Boolean }\n   bunx prisma db push\n\n۲. API اضافه کن (content route.ts):\n   models.project = db.project\n\n۳. بخش تو page.tsx:\n   <section id='projects'>...</section>\n   <ArchiveGrid items={siteContent.projects} type='book' />\n\n۴. لینک تو navbar:\n   <a href='#projects'>projects</a>")

# CH10
doc.add_heading("فصل ۱۰: تم‌ساز رنگی", level=1)
rtl("از پنل ادمین تب themes می‌تونی تم بسازی.")
rtl("۱۵ رنگ قابل تنظیم: bg, primary, primaryBright, text, border, accent, red, cyan و غیره")
rtl("مثال: ساخت تم قرمز:")
code("پنل themes + new theme\nname: Red Alert\nbg: #1a0000\nprimary: #ff3333\nprimaryBright: #ff5555\ntext: #ffcccc\nborder: #3a1010\nSave -> تم تو navbar ظاهر می‌شه!")

# CH11
doc.add_heading("فصل ۱۱: ایمیل و Bale", level=1)
rtl("ایمیل (ساده — یه کادر):")
code("پنل settings -> Email Forwarding\nایمیلت رو وارد کن: myname@gmail.com\nSave\n\nوقتی بازدیدکننده فرم تماس پر می‌کنه:\n۱. پیام توی DB ذخیره می‌شه\n۲. کپی به ایمیلت فرستاده می‌شه (formsubmit.co)")
rtl("Bale:")
code("۱. Bale -> @botfather -> /newbot\n۲. Token کپی کن\n۳. Chat ID بگیر:\n   curl https://api.bale.ai/v1/botsTOKEN/getUpdates\n۴. پنل settings -> Bale -> Token و Chat ID\n۵. Webhook:\n   curl https://api.bale.ai/v1/botsTOKEN/setWebhook?url=DOMAIN/api/bale/webhook\n\nدستورات: /list /reply /disable /enable /stats")

# CH12
doc.add_heading("فصل ۱۲: ضد کپی", level=1)
rtl("۳ لایه محافظت:")
bullet("قفل دامنه: فقط روی دامنه‌های مجاز کار می‌کنه")
bullet("واترمارک نامرئی روی Canvas: شناسه توی پیکسل‌ها مخفی")
bullet("راست‌کلیک و DevTools غیرفعال")
code("// قفل دامنه:\nif (!authorizedDomains.includes(hostname)) {\n  document.body.innerHTML = 'UNAUTHORIZED COPY';\n}\n\n// واترمارک: LSB steganography\n// شناسه: EHSANMORAD-V19-2026")

# CH13
doc.add_heading("فصل ۱۳: دستورات Docker", level=1)
code("docker-compose up -d        # اجرا\ndocker-compose down         # توقف\ndocker-compose logs -f       # لاگ\ndocker-compose up -d --build # بازسازی\ndocker-compose ps            # وضعیت")
rtl("بک‌آپ:")
code("docker cp personal-site:/app/db/custom.db ./backup.db")
rtl("انتقال به VPS:")
code("unzip personal-site-v7.zip\ncd personal-site-v7\ndocker-compose up -d\n# تمام! دامنه‌ت رو به AUTHORIZED_DOMAINS اضافه کن")

# CH14
doc.add_heading("فصل ۱۴: رفع اشکال", level=1)
rtl("سایت باز نمی‌شه:")
code("docker-compose ps     # وضعیت\ndocker-compose logs -f  # لاگ")
rtl("پنل ادمین:")
bullet("آدرس: #admin یا Ctrl+Shift+A")
bullet("رمز: admin123")
rtl("چت AI جواب نمی‌ده:")
bullet("settings -> kill switch بررسی کن")
bullet("settings -> AI Providers بررسی کن")
rtl("آموزش لود نمی‌شه — URL امبد رو چک کن:")
code("درست: https://www.aparat.com/video/video/embed/videohash/XXXX/vframe\nغلط: https://www.aparat.com/v/XXXX")

# SUMMARY
doc.add_heading("خلاصه", level=1)
summary = [
    ("اجرای سایت", "docker-compose up -d"),
    ("ورود به پنل", "#admin یا Ctrl+Shift+A"),
    ("تغییر متن", "پنل -> content"),
    ("تغییر رنگ", "پنل -> themes"),
    ("تغییر رمز", "پنل -> settings -> Security"),
    ("اضافه‌کردن آموزش", "پنل -> content -> Tutorials"),
    ("ساخت تم", "پنل -> themes -> new theme"),
    ("تنظیم ایمیل", "پنل -> settings -> Email"),
    ("تنظیم Bale", "پنل -> settings -> Bale"),
    ("پاک‌کردن چت", "پنل -> settings -> Danger Zone"),
    ("بک‌آپ", "docker cp ... db/custom.db"),
    ("انتقال VPS", "unzip -> docker-compose up -d"),
]
tbl2 = doc.add_table(rows=len(summary)+1, cols=2)
tbl2.style = "Table Grid"
tbl2.rows[0].cells[0].text = "کار"
tbl2.rows[0].cells[1].text = "روش"
for i, (a, b) in enumerate(summary):
    tbl2.rows[i+1].cells[0].text = a
    tbl2.rows[i+1].cells[1].text = b

doc.save(OUTPUT)
print("ذخیره شد: " + OUTPUT)
print("حجم: " + str(os.path.getsize(OUTPUT)) + " bytes")
